#!/usr/bin/env python3
"""quiz-builder harness. Builds a stand-in template with the official layout's features (instruction
rows, merged cells, header on row 8, time dropdown) because Kahoot's own file is not redistributed,
then checks: refusal without a template, a clean build and read back of the example, every planted
rule failure in a bad spec, numeric ordering, placement, and the host key.
Usage: python3 evals/quiz_builder_check.py"""
import copy, json, os, subprocess, sys, tempfile
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SK = os.path.join(ROOT, "plugins", "officer-kit", "skills", "quiz-builder")
SC = os.path.join(SK, "scripts")
SRC = os.path.join(ROOT, "plugins", "officer-kit", "skills", "study-guide", "references", "example-source.txt")
sys.path.insert(0, SC)
import quiz_lib as Q
fails, n = [], 0
def ok(c, m):
    global n
    n += 1
    if not c: fails.append(m)
def run(*a):
    r = subprocess.run([sys.executable] + list(a), capture_output=True, text=True)
    return r.returncode, r.stdout + r.stderr
tmp = tempfile.mkdtemp()
tpl = os.path.join(tmp, "StandInTemplate.xlsx")
wb = Workbook(); ws = wb.active
ws["B2"] = "Quiz template (stand in for tests)"; ws["B3"] = "instructions"; ws.merge_cells("B3:H3")
ws.append([]); ws.append([]); ws.append([]); ws.append([])
hdr = ["", "Question - max 120 characters", "Answer 1", "Answer 2", "Answer 3", "Answer 4", "Time limit (sec)", "Correct answer(s)"]
for c, v in enumerate(hdr, 1): ws.cell(8, c).value = v
ws.cell(9, 1).value = 1; ws.cell(9, 2).value = "Example question"; ws.cell(9, 3).value = "A"; ws.cell(9, 4).value = "B"; ws.cell(9, 7).value = 60; ws.cell(9, 8).value = "1,2"
for r in range(10, 30): ws.cell(r, 1).value = r - 8
dv = DataValidation(type="list", formula1='"5,10,20,30,60,120"'); dv.add("G9"); ws.add_data_validation(dv)
wb.save(tpl)
good = os.path.join(SK, "references", "example-spec.json")
out = os.path.join(tmp, "out")
code, o = run(os.path.join(SC, "build_kahoot.py"), good, "--out", out)
ok(code == 1 and "NOT BUILT: no template" in o, "must refuse without a template")
code, o = run(os.path.join(SC, "build_kahoot.py"), good, "--template", tpl, "--out", out, "--source", SRC, "--report", os.path.join(out, "gates.md"))
ok(code == 0 and "RESULT: CLEAN" in o, "example should build clean:\n" + o)
ok(o.count("READ BACK") == 2 and "FAIL" not in o, "both decks should read back clean")
ok(os.path.exists(os.path.join(out, "gates.md")), "gate report not written")
f1 = os.path.join(out, "Sourdough 1 - Starter and feeding.xlsx")
if os.path.exists(f1):
    w = load_workbook(f1).active
    ok(w["B3"].value == "instructions" and any(str(r) == "B3:H3" for r in w.merged_cells.ranges), "template rows or merges lost")
    ok(any(str(d.sqref) == "G9:G16" for d in w.data_validations.dataValidation), "time dropdown not extended to every question")
    ok(w.cell(9, 2).value != "Example question", "example row not overwritten")
    ok(w.cell(9, 8).value in ("1", "2", "3", "4"), "slot not written as a string number")
    rows = [[w.cell(r, c).value for c in range(2, 9)] for r in range(9, 17)]
    num = [r for r in rows if r[0].startswith("How many hours")][0]
    ok(num[1:5] == ["2", "4", "8", "12"] and num[6] == "2", "numeric options should be ascending with the slot following the value")
s = Q.place(Q.load(good))
for d in s["decks"]:
    seq = [q["a"] for q in d["questions"]]
    ok(all(seq[i] != seq[i - 1] for i in range(1, len(seq))), "adjacent slots repeat in " + d["name"])
    ok(all(q["o"][q["a"]] == q["_intended"] for q in d["questions"]), "placement moved an answer in " + d["name"])
bad = copy.deepcopy(Q.load(good))
qs = bad["decks"][0]["questions"]
qs[0]["q"] = "x" * 130
qs[1]["o"][3] = "All of the above"
qs[2]["q"] = "Which of these is not a sign of a hungry starter?"
qs[4]["o"][2] = ""
qs[5]["o"][0] = "When the moon is full"; qs[5]["a"] = 0; qs[5].pop("src", None)
qs[6]["time"] = 45
qs[7]["q"] = qs[3]["q"]
long_deck = {"name": "Length cue", "questions": []}
for i in range(10):
    long_deck["questions"].append({"q": "Length cue question %d about the float test?" % i, "o": ["If it floats, the starter is ready to bake with", "Hungry", "Overdue", "Sluggish"], "a": 0, "e": "x", "d": "medium", "src": "ready to bake with"})
bad["decks"].append(long_deck)
bp = os.path.join(tmp, "bad.json"); json.dump(bad, open(bp, "w"))
code, o = run(os.path.join(SC, "build_kahoot.py"), bp, "--template", tpl, "--out", os.path.join(tmp, "bad"), "--source", SRC)
ok(code == 1 and "NOT BUILT" in o, "bad spec must not build")
for needle, why in (("characters, limit 120", "long question"), ("of the above", "all of the above"), ("negative stem", "lower case not"),
                    ("an option is empty", "empty option"), ("not found in the source", "untraceable answer"), ("time 45", "time not in template"),
                    ("duplicate", "duplicate question"), ("longest option", "length cue")):
    ok(needle in o, "bad spec should fail on " + why)
ok(not os.path.exists(os.path.join(tmp, "bad")) or not os.listdir(os.path.join(tmp, "bad")), "nothing may be written when rules fail")
stair = {"decks": [{"name": "Stair", "questions": [
    {"q": "Stair question %d?" % i, "o": ["alpha %d" % i, "bravo %d" % i, "charlie %d" % i, "delta %d" % i], "a": i % 4, "e": "x", "d": "medium"} for i in range(6)]}]}
f, w = Q.check(stair)
ok(any("next tile" in x or "cycle" in x for x in f), "a 1,2,3,4 answer staircase should fail")
cue = {"decks": [{"name": "Cue", "questions": [
    {"q": "In a learner centric model, what is the main focus of effort?", "o": ["The learner centric focus", "The instructor", "The commander", "The schoolhouse"], "a": 0, "e": "x", "d": "easy"}]}]}
f, w = Q.check(cue)
ok(any("repeats a stem word" in x for x in w), "a stem that gives away the answer should warn")
long_q = {"q": "x" * 110, "o": ["y" * 60, "z" * 60, "w" * 20], "a": 0, "d": "easy"}
ok(Q.time_for(long_q) == 30, "a long question should get at least 30 seconds")
code, o = run(os.path.join(SC, "build_host_key.py"), good, os.path.join(tmp, "key.docx"))
ok(code == 0 and os.path.exists(os.path.join(tmp, "key.docx")), "host key should build")
print("QUIZ BUILDER CHECK: %d items" % n)
for f in fails: print("  FAIL ", f)
print("-> %d of %d pass" % (n - len(fails), n))
sys.exit(1 if fails else 0)
