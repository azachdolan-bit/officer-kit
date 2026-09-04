#!/usr/bin/env python3
"""Gate 2, mechanical QC for a naval letter (.docx). Deterministic; no judgment calls.

Checks the things that are cheap to get wrong and expensive to send out wrong:
required heading elements, reference integrity, citation resolution both ways,
enclosure naming, signature block spacing, prohibited characters, paragraph
sequence, and placeholders left from an incomplete identity block.

Usage:  python3 qc_letter.py "<letter>.docx" [--encl-title "exact title of enclosure 1"] ...
Exit 0 = clean (warnings allowed), 1 = defects found.
"""
import re
import sys
import unicodedata
from collections import Counter

try:
    import docx
except ImportError:
    sys.exit("python-docx is required: pip install python-docx")

FAIL, WARN = [], []


def fail(m):
    FAIL.append(m)


def warn(m):
    WARN.append(m)


def paras(path):
    d = docx.Document(path)
    return [p.text.rstrip() for p in d.paragraphs]


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    path = sys.argv[1]
    encl_titles = []
    a = sys.argv[2:]
    while a:
        if a[0] == "--encl-title" and len(a) > 1:
            encl_titles.append(a[1])
            a = a[2:]
        else:
            a = a[1:]

    P = paras(path)
    body = "\n".join(P)

    # ---- 0. placeholders from an incomplete identity block ------------------
    for ph in sorted(set(re.findall(r"\[[A-Z][A-Z ]+\]", body))):
        fail(f"placeholder {ph} left in the letter; fill it from the rules file or the spec")

    # ---- 1. required heading elements ---------------------------------------
    for tag in ("From:", "To:", "Subj:"):
        if not any(p.startswith(tag) for p in P):
            fail(f"missing required heading element {tag}")
    if any(p.strip() == "Via:" for p in P):
        fail("empty Via: line; omit Via entirely when there is no routing requirement")

    subj = next((p for p in P if p.startswith("Subj:")), "")
    subj_txt = subj[5:].strip()
    if subj_txt and subj_txt != subj_txt.upper():
        fail("Subj line is not in all caps")
    if subj_txt.endswith("."):
        warn("Subj line ends with a period; it is a phrase, not a sentence")

    # ---- 2. reference list --------------------------------------------------
    ref_start = next((i for i, p in enumerate(P) if p.startswith("Ref:")), None)
    refs = {}
    if ref_start is not None:
        i = ref_start
        while i < len(P):
            m = re.match(r"(?:Ref:\s*)?\(([a-z])\)\s*(.+)", P[i].strip())
            if not m:
                if i > ref_start:
                    break
            else:
                refs[m.group(1)] = m.group(2).strip()
            i += 1

    if refs:
        letters = list(refs)
        expect = [chr(ord("a") + n) for n in range(len(letters))]
        if letters != expect:
            fail(f"reference letters not sequential: listed {letters}, expected {expect}")
        norm = [re.sub(r"[^a-z0-9]", "", t.lower()) for t in refs.values()]
        for t, n in Counter(norm).items():
            if n > 1:
                dup = [k for k, v in refs.items() if re.sub(r"[^a-z0-9]", "", v.lower()) == t]
                fail(f"duplicate reference: {dup} name the same document")
        for k, v in refs.items():
            if re.search(r"\band\b", v) and v.count("(") <= 1:
                warn(f"reference ({k}) may combine two documents: {v!r}")

    # ---- 3. enclosure list --------------------------------------------------
    enc_start = next((i for i, p in enumerate(P) if p.startswith("Encl:")), None)
    encls = {}
    if enc_start is not None:
        i = enc_start
        while i < len(P):
            m = re.match(r"(?:Encl:\s*)?\((\d+)\)\s*(.+)", P[i].strip())
            if not m:
                if i > enc_start:
                    break
            else:
                encls[m.group(1)] = m.group(2).strip()
            i += 1
    for n, want in enumerate(encl_titles, start=1):
        got = encls.get(str(n))
        if got is None:
            fail(f"enclosure ({n}) declared on the command line but not listed in the letter")
        elif got.strip().lower() != want.strip().lower():
            fail(f"enclosure ({n}) title mismatch\n     listed: {got!r}\n     actual: {want!r}")

    # ---- 4. citations resolve both ways -------------------------------------
    head_end = max(x for x in (enc_start, ref_start, 0) if x is not None)
    text_after_head = "\n".join(P[head_end + 1:])
    cited = set()
    for m in re.finditer(r"references?\s*\(([a-z])\)\s*(?:through|thru|to)\s*\(([a-z])\)", text_after_head, re.I):
        for c in range(ord(m.group(1)), ord(m.group(2)) + 1):
            cited.add(chr(c))
    for m in re.finditer(r"references?\s*\(([a-z])\)", text_after_head, re.I):
        cited.add(m.group(1))
    enc_cited = set(re.findall(r"enclosures?\s*\((\d+)\)", text_after_head, re.I))

    for k in refs:
        if k not in cited:
            fail(f"reference ({k}) is listed but never cited in the body")
    for k in sorted(cited):
        if k not in refs:
            fail(f"body cites reference ({k}) which is not in the reference list")
    for k in encls:
        if k not in enc_cited:
            warn(f"enclosure ({k}) is listed but never mentioned in the body")
    for k in sorted(enc_cited):
        if k not in encls:
            fail(f"body cites enclosure ({k}) which is not in the enclosure list")

    # ---- 5. signature block: fourth line below the text ---------------------
    sig_i = None
    for i in range(len(P) - 1, -1, -1):
        s = P[i].strip()
        if s and re.fullmatch(r"[A-Z][A-Z.\s]{3,}", s) and "." in s:
            sig_i = i
            break
    if sig_i is None:
        warn("no signature block found (expected an ALL CAPS name line with initials)")
    else:
        prev = max((i for i in range(sig_i) if P[i].strip()), default=None)
        gap = sig_i - prev - 1
        if gap != 3:
            fail(f"signature block sits {gap + 1} lines below the text; the standard is the fourth line "
                 f"(three blank lines). Cut body text to fit a page, never the signature spacing.")

    # ---- 6. prohibited characters and typography ----------------------------
    for i, p in enumerate(P):
        if "—" in p or "–" in p:
            fail(f"em or en dash in paragraph {i + 1}: {p[:70]!r}")
        if re.search(r"[a-z][a-z]  +[a-z]", p):
            warn(f"double space inside a sentence, paragraph {i + 1}: {p[:60]!r}")
        if " : " in p:
            warn(f"stray ' : ' in paragraph {i + 1}; rewrite the sentence rather than swapping a dash for a colon")
    for ch in set(body):
        if unicodedata.category(ch) == "Cc" and ch not in "\n\t":
            fail(f"control character U+{ord(ch):04X} in the text")
    for w in ("colour", "organisation", "organised", "centre", "recognise", "behaviour"):
        if re.search(rf"\b{w}\b", body, re.I):
            warn(f"British spelling {w!r} in the text")

    # ---- 7. numbered paragraphs sequential ----------------------------------
    nums = [int(m.group(1)) for p in P if (m := re.match(r"(\d+)\.\s", p.strip()))]
    if nums and nums != list(range(1, len(nums) + 1)):
        fail(f"numbered paragraphs out of sequence: {nums}")

    # ---- report ---------------------------------------------------------------
    print(f"QC: {path}")
    print(f"  references: {len(refs)}  enclosures: {len(encls)}  numbered paragraphs: {len(nums)}")
    for m in WARN:
        print(f"  WARN  {m}")
    for m in FAIL:
        print(f"  FAIL  {m}")
    if not FAIL and not WARN:
        print("  clean")
    print(f"  -> {len(FAIL)} failures, {len(WARN)} warnings")
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
