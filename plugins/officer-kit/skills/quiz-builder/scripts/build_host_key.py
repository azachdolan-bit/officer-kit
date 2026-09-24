#!/usr/bin/env python3
"""Host answer and page key for a set of decks: for every question, the correct answer, why, and
where the source teaches it, so the host can send a player to the page. Landscape .docx.
Usage: python3 build_host_key.py <spec.json> <out.docx>   (run after build_kahoot.py on the same spec)"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import quiz_lib as Q
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.shared import Pt, Inches

def main(spec_path, out):
    spec = Q.place(Q.load(spec_path))
    doc = Document()
    sec = doc.sections[0]; sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for m in ("left_margin", "right_margin"): setattr(sec, m, Inches(0.6))
    doc.styles["Normal"].font.size = Pt(9.5)
    doc.add_heading(spec.get("title", "Quiz") + ": host answer and page key", level=0)
    for d in spec["decks"]:
        doc.add_heading(d["name"], level=1)
        t = doc.add_table(rows=1, cols=5); t.style = "Table Grid"
        for c, h in zip(t.rows[0].cells, ("No.", "Question", "Answer", "Why", "Where it is")):
            c.text = h; c.paragraphs[0].runs[0].bold = True
        t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for i, q in enumerate(d["questions"], 1):
            row = t.add_row()
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            vals = (str(i), q["q"], "%d. %s" % (q["a"] + 1, q["o"][q["a"]]), q["e"], q.get("ref", "not given"))
            for c, v in zip(row.cells, vals): c.text = v
        for row in t.rows:
            for c, w in zip(row.cells, (0.4, 3.2, 2.2, 2.6, 1.6)): c.width = Inches(w)
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    doc.save(out); print("written", out)
    return 0

if __name__ == "__main__":
    if len(sys.argv) != 3: sys.exit(__doc__)
    sys.exit(main(*sys.argv[1:]))
