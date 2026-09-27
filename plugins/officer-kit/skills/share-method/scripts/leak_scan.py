#!/usr/bin/env python3
"""Leak scan for a file meant to be shared: does it carry any content from the sources it was built from?

Usage: python3 leak_scan.py <shared file> --source <file> [<file> ...] [--product <file> ...] [--terms terms.txt]
                            [--allow allow.txt] [--k 8] [--report scan.md]

Four passes, every hit shown with its line so it can be judged in context:
  1. Terms: acronyms, code and model numbers, and capitalized names (matched as capitalized) taken from the sources, plus the
     user's own terms file. Whole word matching only; an all capitals term matches case sensitively,
     so "not" never matches OT and "mistakes" never matches MIST. Substring only matches are counted
     and listed as not leaks, so the reader sees they were considered.
  2. Values: every number of 13 or more (except 100) that appears in both the shared file and a source.
  3. Copied runs: any run of k or more words (default 8) that the shared file shares with a source,
     after case, punctuation, and whitespace are normalized. A run shared only with one of the user's
     own products (--product) is the user's method wording unless it carries a term or a value, so it
     is listed for one read rather than failed.
  4. Mechanism: any mention of browser automation. This one cannot be allowed.
  Then, for one read and never failing on its own: every word of five letters or more that the shared
  file shares with a source and that the kit's own method text (SKILL.md and the references) never
  uses, and every number from 2 to 12 shared with a source. A subject word in lower case, or a small
  threshold, shows up here.
A hit on a term listed in the allow file ("term | reason", reason required) is reported as allowed.
Exit 0 when clean, 1 when anything real is left, 2 on a usage error.
Text sources: .md .txt .html .csv; .docx and .pdf are read when python-docx or pdftotext is present."""
import argparse, os, re, subprocess, sys, zipfile

STOP_CAPS = set("""A AN AND ARE AS AT BE BY DO FOR FROM HE HOW I IF IN IS IT ITS NO NOT OF ON OR SO THE THEN
THIS TO UP US WE WHAT WHEN WHO WHY WITH YOU YOUR ALL ANY ONE TWO THREE FOUR FIVE SIX SEVEN EIGHT NINE TEN
NOTE NOTES WARNING CAUTION IMPORTANT STEP STEPS LINE LINES PAGE PAGES PART SECTION TABLE FIGURE ANNEX
APPENDIX CHAPTER END START STOP YES OK NEW OLD TOP SEE ALSO EACH EVERY ONLY MUST NEVER ALWAYS DONE PASS
FAIL TRUE FALSE II III IV VI VII VIII IX XI XII AM PM USA US UK ID FAQ PDF HTML MD CSV DOCX XLSX JSON URL
INSTRUCTIONS CLAUDE AI OR NA N/A TBD""".split())
STOP_LEAD = set("the a an this that these those step line page part section table figure note see and or of for to in on at by with from each every if when".split())
MECHANISM = [r"browser automation", r"claude in chrome", r"chrome extension", r"playwright", r"selenium", r"puppeteer",
             r"headless", r"computer use", r"drive (?:the|a|your) browser", r"automate[ds]? (?:the|a|your) browser",
             r"browser (?:agent|tool|extension)"]
COMMON = set("""about above after again against along among another around because become before being below
between beyond block build built cannot check clear close could count doing during early either every first
follow found given going great group heard until later least leave level might never night number often other
place point quick quite ready right round second seven shall short should since small sound start state still
their there these thing think those three through today under until using usual value where which while whole
whose would write wrong years""".split())
_HERE = os.path.dirname(os.path.abspath(__file__))
def _kit_text():
    out = []
    for rel in ("../SKILL.md", "../references/builder-template.md", "../references/method-inventory.md"):
        try:
            out.append(open(os.path.join(_HERE, rel), encoding="utf-8").read())
        except OSError:
            pass
    return "\n".join(out)
KIT_TEXT = _kit_text()
NUM = re.compile(r"(?<![\w.])\d{1,3}(?:,\d{3})+(?:\.\d+)?|(?<![\w.])\d+(?:\.\d+)?(?![\w])")

def read(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        try:
            import docx
            d = docx.Document(path)
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for r in t.rows:
                    parts.append(" | ".join(c.text for c in r.cells))
            return "\n".join(parts)
        except ImportError:
            x = zipfile.ZipFile(path).read("word/document.xml").decode("utf-8", "ignore")
            return re.sub(r"<[^>]+>", " ", x.replace("</w:p>", "\n"))
    if ext == ".pdf":
        r = subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True)
        if r.returncode:
            sys.exit("cannot read %s: install pdftotext or capture it to text first" % path)
        return r.stdout
    return open(path, encoding="utf-8", errors="ignore").read()

def wordrx(term, kind=""):
    body = re.escape(term).replace(r"\ ", r"\s+")
    exact = kind == "name" or (term.upper() == term and re.search("[A-Z]", term))
    flags = 0 if exact else re.I
    return re.compile(r"(?<![A-Za-z0-9])" + body + r"(?![A-Za-z0-9])", flags)

def source_terms(text):
    """Acronyms (an all capitals word that the source also writes in lower case is an ordinary word
    set in capitals, not an acronym, and is dropped), code and model numbers, capitalized names."""
    terms = {}
    lower_words = set(re.findall(r"(?<![A-Za-z])[a-z]+(?![A-Za-z])", text))
    for m in re.finditer(r"(?<![A-Za-z0-9])[A-Z][A-Z0-9&]{1,}(?:[-/][A-Z0-9]+)*(?![A-Za-z0-9])", text):
        t = m.group(0)
        if t not in STOP_CAPS and not t.isdigit() and t.lower() not in lower_words:
            terms.setdefault(t, "acronym")
    for m in re.finditer(r"(?<![A-Za-z0-9])(?=[A-Za-z0-9/-]*\d)(?=[A-Za-z0-9/-]*[A-Za-z])[A-Za-z0-9]+(?:[-/][A-Za-z0-9]+)+(?![A-Za-z0-9])|(?<![A-Za-z0-9])(?=\w*\d)(?=\w*[A-Z])[A-Z0-9]{3,}(?![A-Za-z0-9])", text):
        terms[m.group(0)] = "code"
    for m in re.finditer(r"(?<![A-Za-z])((?:[A-Z][a-z]+)(?: (?:of |the |and )?[A-Z][a-z]+){1,3})", text):
        t = re.sub(r"\s+", " ", m.group(1))
        if t.split()[0].lower() not in STOP_LEAD:
            terms.setdefault(t, "name")
    return terms

def norm_words(text):
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).split()

def line_of(text, pos):
    ln = text.count("\n", 0, pos) + 1
    s = max(text.rfind("\n", 0, pos) + 1, pos - 70)
    e = text.find("\n", pos)
    e = min(e if e >= 0 else len(text), pos + 90)
    return ln, ("..." if s > text.rfind("\n", 0, pos) + 1 else "") + text[s:e].strip()

def scan(shared, sources, user_terms=(), allow=None, k=8, products=()):
    """sources: the course or unit material. products: the user's own products built from it, whose
    wording is partly the user's method; a run shared only with a product fails only if it carries a
    term or a value from the sources or the product."""
    allow = allow or {}
    src = "\n".join(list(sources) + list(products))
    terms = source_terms(src)
    for t in user_terms:
        terms[t] = "user term"
    hits, allowed, substr, mech = [], [], [], []
    for t, kind in sorted(terms.items()):
        rx = wordrx(t, kind)
        found = list(rx.finditer(shared))
        if not found:
            if len(t) >= 2 and t.lower() in shared.lower():
                substr.append(t)
            continue
        for m in found:
            ln, ctx = line_of(shared, m.start())
            row = (t, kind, ln, ctx)
            (allowed if t in allow else hits).append(row)
    src_nums = set(n.replace(",", "") for n in NUM.findall(src))
    for m in NUM.finditer(shared):
        v = m.group(0).replace(",", "")
        try:
            f = float(v)
        except ValueError:
            continue
        if f < 13 or f == 100 or v not in src_nums:
            continue
        ln, ctx = line_of(shared, m.start())
        (allowed if v in allow else hits).append((v, "value", ln, ctx))
    bw = norm_words(shared)
    term_words, name_phrases = set(), []
    for t, kind in terms.items():
        if kind == "name":
            name_phrases.append(" " + " ".join(norm_words(t)) + " ")
        else:
            term_words.update(norm_words(t))
    def copied(texts):
        ws = norm_words("\n".join(texts))
        grams = set(tuple(ws[i:i + k]) for i in range(len(ws) - k + 1))
        out, i = [], 0
        while i <= len(bw) - k:
            if tuple(bw[i:i + k]) in grams:
                j = i + 1
                while j <= len(bw) - k and tuple(bw[j:j + k]) in grams:
                    j += 1
                out.append(" ".join(bw[i:j + k - 1]))
                i = j + k - 1
            else:
                i += 1
        return out
    runs = copied(sources)
    method_runs = []
    for r in copied(products):
        if r in runs:
            continue
        words = r.split()
        padded = " " + r + " "
        if (any(w in src_nums for w in words) or any(w in term_words for w in words if len(w) > 1)
                or any(ph in padded for ph in name_phrases)):
            runs.append(r)
        else:
            method_runs.append(r)
    runs_real = [r for r in runs if r not in allow]
    runs_allowed = [r for r in runs if r in allow]
    kit_words = set(norm_words(KIT_TEXT)) | COMMON
    src_words = set(w for w in norm_words(src) if len(w) >= 5 and not w.isdigit())
    vocab = sorted(set(w for w in norm_words(shared) if w in src_words and w not in kit_words))
    def list_marker(m):
        before = shared[shared.rfind("\n", 0, m.start()) + 1:m.start()]
        return (not before.strip() and shared[m.end():m.end() + 1] in ".)") or re.search(r"step\s*$", before, re.I)
    small = sorted(set(m.group(0) for m in NUM.finditer(shared) if m.group(0) in src_nums
                       and 2 <= float(m.group(0).replace(",", "")) < 13 and not list_marker(m)), key=float)
    for pat in MECHANISM:
        for m in re.finditer(pat, shared, re.I):
            ln, ctx = line_of(shared, m.start())
            mech.append((m.group(0), ln, ctx))
    return {"terms": terms, "hits": hits, "allowed": allowed, "substring_only": sorted(set(substr)),
            "runs": runs_real, "runs_allowed": runs_allowed, "method_runs": method_runs, "vocab": vocab, "small": small, "mechanism": mech,
            "clean": not hits and not runs_real and not mech}

def load_allow(path):
    out = {}
    for n, line in enumerate(open(path, encoding="utf-8"), 1):
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        if "|" not in line or not line.split("|", 1)[1].strip():
            sys.exit("allow file line %d has no reason: %r (write: term | why it is not a leak)" % (n, line))
        t, why = [x.strip() for x in line.split("|", 1)]
        out[t] = why
    return out

def render(name, r, allow):
    kinds = {}
    for kind in r["terms"].values():
        kinds[kind] = kinds.get(kind, 0) + 1
    L = ["# Leak scan: %s" % name, "",
         "Verdict: %s" % ("CLEAN" if r["clean"] else "NOT CLEAN"), "",
         "Terms checked: %d (%s)" % (len(r["terms"]), ", ".join("%d %s" % (v, k) for k, v in sorted(kinds.items()))), ""]
    L.append("## Term and value hits")
    L += ["%d. line %d, %s `%s`: %s" % (i + 1, ln, kind, t, ctx) for i, (t, kind, ln, ctx) in enumerate(r["hits"])] or ["none"]
    L += ["", "## Copied runs (shared word runs with a source)"]
    L += ["%d. %s" % (i + 1, x) for i, x in enumerate(r["runs"])] or ["none"]
    L += ["", "## Wording shared only with the product, carrying no term or value (method wording; read once)"]
    L += ["%d. %s" % (i + 1, x) for i, x in enumerate(r["method_runs"])] or ["none"]
    L += ["", "## Read once: words and small numbers shared with the sources that the kit's own method text never uses"]
    L += ["Words: " + (", ".join(r["vocab"]) or "none"), "Numbers 2 to 12: " + (", ".join(r["small"]) or "none")]
    L += ["", "## Mechanism mentions (never allowed)"]
    L += ["%d. line %d `%s`: %s" % (i + 1, ln, t, ctx) for i, (t, ln, ctx) in enumerate(r["mechanism"])] or ["none"]
    L += ["", "## Allowed, with the reason given"]
    L += ["%d. `%s`: %s" % (i + 1, t, allow.get(t, "")) for i, t in enumerate(sorted(set([a[0] for a in r["allowed"]] + r["runs_allowed"])))] or ["none"]
    L += ["", "## Substring only matches (inside another word; not leaks)"]
    so = r["substring_only"]
    L += [("%d: " % len(so) + ", ".join(so[:15]) + (" and %d more" % (len(so) - 15) if len(so) > 15 else "")) if so else "none"]
    return "\n".join(L) + "\n"

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("shared")
    ap.add_argument("--source", nargs="+", required=True, help="the course or unit material")
    ap.add_argument("--product", nargs="*", default=[], help="the user's own products built from it")
    ap.add_argument("--terms")
    ap.add_argument("--allow")
    ap.add_argument("--k", type=int, default=8)
    ap.add_argument("--report")
    a = ap.parse_args()
    user_terms = []
    if a.terms:
        user_terms = [l.strip() for l in open(a.terms, encoding="utf-8") if l.strip() and not l.startswith("#")]
    allow = load_allow(a.allow) if a.allow else {}
    r = scan(read(a.shared), [read(s) for s in a.source], user_terms, allow, a.k, [read(p) for p in a.product])
    out = render(os.path.basename(a.shared), r, allow)
    print(out)
    if a.report:
        open(a.report, "w", encoding="utf-8").write(out)
    sys.exit(0 if r["clean"] else 1)

if __name__ == "__main__":
    main()
